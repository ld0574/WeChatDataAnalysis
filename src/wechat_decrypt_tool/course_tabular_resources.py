"""Read-only, hash-bound XLSX cell indexing. Never executes formulas or exports books."""

from __future__ import annotations

import datetime
import hashlib
import io
import json
import time
import zipfile
from pathlib import Path

from .course_attachments import (
    archive_locator,
    digest,
    normalized_original,
    open_archive_member,
)


def workbook_chunks(raw, checkpoint=lambda: None):
    import openpyxl

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if sum(x.file_size for x in entries) > 128 * 1024**2 or any(
            x.file_size > 64 * 1024**2 for x in entries
        ):
            raise ValueError("Spreadsheet expanded-byte bound exceeded; raw retained")
        visuals = sum(
            x.filename.startswith(("xl/media/", "xl/charts/")) and not x.is_dir()
            for x in entries
        )
    book = openpyxl.load_workbook(
        io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
    )
    cache = openpyxl.load_workbook(
        io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
    )
    chunks, sheets = [], []
    count = formulas = missing = total_length = 0

    def scalar(value):
        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
            return value.isoformat()
        if isinstance(value, datetime.timedelta):
            return str(value)
        if value is None or type(value) in (str, bool, int, float):
            return value
        raise ValueError("Unsupported spreadsheet cell value; raw retained")

    try:
        for sheet in book.worksheets:
            if (sheet.max_row or 0) > 200000 or (sheet.max_column or 0) > 512:
                raise ValueError("Spreadsheet dimension bound exceeded; raw retained")
            lines, length, part, sheet_cells = [], 0, 0, 0
            for row, cached in zip(sheet.iter_rows(), cache[sheet.title].iter_rows()):
                checkpoint()
                for cell, saved in zip(row, cached):
                    if cell.value is None:
                        continue
                    count += 1
                    sheet_cells += 1
                    if count > 500000:
                        raise ValueError(
                            "Spreadsheet cell bound exceeded; raw retained"
                        )
                    record = {
                        "cell": cell.coordinate,
                        "type": cell.data_type,
                        "value": scalar(cell.value),
                    }
                    if cell.data_type == "f":
                        formulas += 1
                        missing += saved.value is None
                        record.update(
                            cached_value=scalar(saved.value), recalculated=False
                        )
                    text = json.dumps(record, ensure_ascii=False, allow_nan=False)
                    total_length += len(text) + 1
                    if total_length > 32 * 1024**2:
                        raise ValueError(
                            "Spreadsheet text bound exceeded; raw retained"
                        )
                    if lines and length + len(text) > 3000:
                        chunks.append(
                            (f"sheet {sheet.title} part {part}", "\n".join(lines))
                        )
                        lines, length, part = [], 0, part + 1
                    lines.append(text)
                    length += len(text) + 1
            if lines:
                chunks.append((f"sheet {sheet.title} part {part}", "\n".join(lines)))
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "nonempty_cells": sheet_cells,
                }
            )
    finally:
        book.close()
        cache.close()
    return chunks, {
        "sheets": sheets,
        "nonempty_cells": count,
        "formula_cells": formulas,
        "formula_cached_values_missing": missing,
        "recalculated": False,
        "visual_assets_raw_retained_not_rendered": visuals,
    }


def index_spreadsheets(store, audit_path):
    import openpyxl

    audit = json.loads(Path(audit_path).read_text(encoding="utf-8"))
    manifest = json.loads(
        (store.root / "resources-manifest.json").read_text(encoding="utf-8")
    )
    resources = {
        normalized_original(manifest["source"] + "/" + r["relative_path"]): r
        for r in manifest["completed"]
    }
    targets, seen = [], set()
    for item in audit["completed"]:
        original = resources[normalized_original(item["path"])]
        for key, resource in resources.items():
            if resource["sha256"] != original["sha256"]:
                continue
            path = str(Path(manifest["source"]) / resource["relative_path"])
            pointer = (key, tuple(item["chain"]))
            if pointer not in seen:
                seen.add(pointer)
                targets.append({**item, "path": path, "resource": resource})
    extractor = "xlsx-cells-v1/openpyxl-" + openpyxl.__version__
    completed, errors, checked, parsed = [], [], set(), {}
    for item in targets:
        store.checkpoint()
        store.write_json(
            "spreadsheet-index-progress.json",
            {
                "stage": "indexing",
                "pid": __import__("os").getpid(),
                "updated": time.time(),
                "completed": len(completed),
                "errors": len(errors),
                "total": len(targets),
            },
        )
        try:
            resource = item["resource"]
            source = (store.root / resource["object"]).resolve()
            if not source.is_relative_to(
                (store.root / "resources" / "objects").resolve()
            ):
                raise ValueError("Spreadsheet object outside retained resource root")
            before = source.stat()
            state = (str(source), before.st_size, before.st_mtime_ns)
            if before.st_size != resource["bytes"] or (
                state not in checked and digest(source) != resource["sha256"]
            ):
                raise RuntimeError("Retained spreadsheet container changed")
            checked.add(state)
            key = archive_locator(item["chain"]) if item["chain"] else ""
            if item["chain"]:
                with open_archive_member(
                    source, item["chain"], checkpoint=store.checkpoint
                ) as (entry, stream):
                    if entry.file_size > 32 * 1024**2:
                        raise ValueError("Spreadsheet input bound exceeded")
                    raw = stream.read(32 * 1024**2 + 1)
                    if len(raw) != entry.file_size:
                        raise ValueError("Spreadsheet member byte count mismatch")
            else:
                if source.stat().st_size > 32 * 1024**2:
                    raise ValueError("Spreadsheet input bound exceeded")
                raw = source.read_bytes()
            sha = hashlib.sha256(raw).hexdigest()
            if sha != item["raw_digest"]:
                raise RuntimeError("Spreadsheet differs from audited revision")
            if store.document_revision(item["path"], key, sha, extractor) is not None:
                completed.append(
                    {"path": item["path"], "source_key": key, "cached": True}
                )
                continue
            if sha not in parsed:
                parsed[sha] = workbook_chunks(raw, store.checkpoint)
            chunks, metadata = parsed[sha]
            after = source.stat()
            if (before.st_size, before.st_mtime_ns) != (
                after.st_size,
                after.st_mtime_ns,
            ):
                raise RuntimeError("Spreadsheet container changed during indexing")
            store.save_document_versions(
                [
                    {
                        "path": item["path"],
                        "source_key": key,
                        "raw_digest": sha,
                        "extractor": extractor,
                        "chunks": [
                            (key + " " + location, text) for location, text in chunks
                        ],
                        "metadata": {
                            **metadata,
                            "archive_members": item["chain"],
                            "raw_source_retained": True,
                        },
                    }
                ]
            )
            completed.append(
                {
                    "path": item["path"],
                    "source_key": key,
                    "nonempty_cells": metadata["nonempty_cells"],
                }
            )
        except InterruptedError:
            raise
        except Exception as exc:
            errors.append(
                {
                    "path": item["path"],
                    "chain": item["chain"],
                    "error": type(exc).__name__ + ": " + str(exc),
                }
            )
        store.write_json(
            "spreadsheet-index-report.json",
            {
                "completed": completed,
                "errors": errors,
                "requested": len(targets),
                "originals_not_modified": True,
            },
        )
    result = {
        "completed": completed,
        "errors": errors,
        "requested": len(targets),
        "originals_not_modified": True,
    }
    store.write_json("spreadsheet-index-report.json", result)
    store.write_json(
        "spreadsheet-index-progress.json",
        {
            "stage": "finished_with_errors" if errors else "batch_finished",
            "pid": __import__("os").getpid(),
            "updated": time.time(),
            "completed": len(completed),
            "errors": len(errors),
            "total": len(targets),
        },
    )
    return result
